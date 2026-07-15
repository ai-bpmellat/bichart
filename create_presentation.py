"""
Generate a Persian presentation (PPTX) for the Rayamate BI project.
Run: python create_presentation.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from copy import deepcopy
from lxml import etree

# Theme colors (Rayamate green)
GREEN = RGBColor(0x05, 0x96, 0x69)
GREEN_DARK = RGBColor(0x04, 0x78, 0x57)
GREEN_SOFT = RGBColor(0xEC, 0xFD, 0xF5)
INK = RGBColor(0x0F, 0x2A, 0x1F)
MUTED = RGBColor(0x5A, 0x7A, 0x6D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF4, 0xF8, 0xF6)
ACCENT_ORANGE = RGBColor(0xD9, 0x77, 0x06)

FONT = "Tahoma"  # Reliable Persian support on Windows


def set_run_font(run, size=18, bold=False, color=INK, font_name=FONT):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # Force East Asian / complex script font for Persian
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:cs", "a:ea"):
        el = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}%s" % tag.split(":")[1])
        # simpler: set via typeface attributes if present
    try:
        from pptx.oxml.ns import qn
        for local in ("latin", "ea", "cs"):
            node = rPr.find(qn(f"a:{local}"))
            if node is None:
                node = etree.SubElement(rPr, qn(f"a:{local}"))
            node.set("typeface", font_name)
    except Exception:
        pass


def add_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_round_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def set_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=18,
    bold=False,
    color=INK,
    align=PP_ALIGN.RIGHT,
    font_name=FONT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, color=color, font_name=font_name)
    return box


def add_bullets(slide, left, top, width, height, items, size=16, color=INK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.RIGHT
        p.level = 0
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = f"•  {item}"
        set_run_font(run, size=size, color=color)
    return box


def footer(slide, page, total=11):
    set_textbox(
        slide,
        Inches(0.4),
        Inches(7.05),
        Inches(9.2),
        Inches(0.3),
        f"Rayamate  |  رایامیت  |  صفحه {page} از {total}",
        size=11,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    total = 11

    # ----- 1 Title -----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, CREAM)
    add_rect(s, 0, 0, Inches(0.35), prs.slide_height, GREEN)
    add_round_rect(s, Inches(8.8), Inches(-1.2), Inches(5.5), Inches(3.2), GREEN_SOFT)
    add_round_rect(s, Inches(-1.0), Inches(5.2), Inches(4.5), Inches(3.0), GREEN_SOFT)
    set_textbox(s, Inches(1.0), Inches(1.8), Inches(11), Inches(0.8), "RAYAMATE", size=42, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
    set_textbox(s, Inches(1.0), Inches(2.55), Inches(11), Inches(0.6), "رایامیت", size=36, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    set_textbox(
        s,
        Inches(1.5),
        Inches(3.4),
        Inches(10.3),
        Inches(0.8),
        "دستیار هوشمند هوش تجاری برای تحلیل گفتگویی تراکنش‌ها",
        size=24,
        bold=True,
        color=INK,
        align=PP_ALIGN.CENTER,
    )
    set_textbox(
        s,
        Inches(2.0),
        Inches(4.3),
        Inches(9.3),
        Inches(0.7),
        "پرسش به زبان طبیعی ← تولید امن SQL ← اجرای گزارش ← تحلیل هوشمند",
        size=16,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    set_textbox(
        s,
        Inches(2.0),
        Inches(5.5),
        Inches(9.3),
        Inches(0.4),
        "ارائه فنی پروژه  |  بهپرداخت ملت",
        size=14,
        color=GREEN_DARK,
        align=PP_ALIGN.CENTER,
    )
    footer(s, 1, total)

    # ----- 2 Problem -----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    add_rect(s, 0, 0, prs.slide_width, Inches(1.1), GREEN)
    set_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.5), "چالش اصلی", size=28, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    add_bullets(
        s,
        Inches(0.8),
        Inches(1.6),
        Inches(11.5),
        Inches(4.8),
        [
            "دسترسی به گزارش‌های تراکنشی معمولاً نیازمند دانش SQL و ابزار BI تخصصی است.",
            "کاربران کسب‌وکار می‌خواهند به زبان فارسی یا انگلیسی سؤال بپرسند و سریع پاسخ بگیرند.",
            "اجرای مستقیم کوئری تولیدشده توسط مدل بدون کنترل امنیتی، خطرناک است.",
            "نیاز به احراز هویت، مدیریت کاربران و ذخیره تاریخچه گفتگو برای هر کاربر وجود دارد.",
            "در محیط عملیاتی باید هم مدل ابری (AvalAI) و هم مدل محلی (Ollama) پشتیبانی شود.",
        ],
        size=18,
    )
    footer(s, 2, total)

    # ----- 3 Solution -----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    add_rect(s, 0, 0, prs.slide_width, Inches(1.1), GREEN)
    set_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.5), "راه‌حل: رایامیت", size=28, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    cards = [
        ("گفتگوی هوشمند", "کاربر سؤال خود را به زبان طبیعی می‌پرسد و سیستم گزارش می‌سازد."),
        ("SQL امن", "فقط کوئری‌های خواندنی (SELECT) پس از اعتبارسنجی امنیتی اجرا می‌شوند."),
        ("تحلیل افزوده", "پس از نمایش نتیجه، کاربر می‌تواند تحلیل و پیش‌بینی هوشمند دریافت کند."),
        ("تجربه یکپارچه", "صفحه معرفی، ورود، داشبورد گفتگو و مدیریت کاربران در یک سامانه."),
    ]
    for i, (title, body) in enumerate(cards):
        col = i % 2
        row = i // 2
        left = Inches(0.7 + col * 6.2)
        top = Inches(1.5 + row * 2.4)
        add_round_rect(s, left, top, Inches(5.8), Inches(2.1), GREEN_SOFT)
        set_textbox(s, left + Inches(0.3), top + Inches(0.35), Inches(5.2), Inches(0.5), title, size=20, bold=True, color=GREEN_DARK, align=PP_ALIGN.RIGHT)
        set_textbox(s, left + Inches(0.3), top + Inches(0.9), Inches(5.2), Inches(0.9), body, size=15, color=INK, align=PP_ALIGN.RIGHT)
    footer(s, 3, total)

    # ----- 4 Architecture -----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    add_rect(s, 0, 0, prs.slide_width, Inches(1.1), GREEN)
    set_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.5), "معماری سامانه", size=28, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    layers = [
        ("رابط کاربری", "صفحه معرفی، ورود با کپچا، داشبورد گفتگو و پنل تاریخچه"),
        ("API (FastAPI)", "مسیرهای /api/chat ، /api/analyze ، /api/users ، /api/history"),
        ("لایه هوش مصنوعی", "Ollama محلی یا AvalAI ابری برای تولید SQL و تحلیل"),
        ("امنیت و داده", "اعتبارسنجی SQL، نشست کاربران، پایگاه SQLite/قابل تعویض"),
    ]
    for i, (title, body) in enumerate(layers):
        top = Inches(1.45 + i * 1.25)
        add_round_rect(s, Inches(1.5), top, Inches(10.3), Inches(1.05), GREEN_SOFT if i % 2 == 0 else CREAM)
        set_textbox(s, Inches(1.8), top + Inches(0.15), Inches(9.7), Inches(0.35), title, size=18, bold=True, color=GREEN_DARK, align=PP_ALIGN.RIGHT)
        set_textbox(s, Inches(1.8), top + Inches(0.5), Inches(9.7), Inches(0.4), body, size=14, color=INK, align=PP_ALIGN.RIGHT)
    footer(s, 4, total)

    # ----- 5 Pipeline -----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    add_rect(s, 0, 0, prs.slide_width, Inches(1.1), GREEN)
    set_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.5), "گردش کار پاسخ به سؤال", size=28, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    steps = [
        ("۱", "دریافت سؤال کاربر"),
        ("۲", "تولید SQL توسط مدل"),
        ("۳", "بررسی ایمنی کوئری"),
        ("۴", "اجرا روی پایگاه داده"),
        ("۵", "نمایش جدول و نمودار"),
        ("۶", "تحلیل اختیاری"),
    ]
    for i, (num, label) in enumerate(steps):
        left = Inches(0.55 + i * 2.15)
        add_round_rect(s, left, Inches(2.4), Inches(2.0), Inches(2.3), GREEN_SOFT)
        set_textbox(s, left, Inches(2.65), Inches(2.0), Inches(0.6), num, size=28, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
        set_textbox(s, left + Inches(0.1), Inches(3.4), Inches(1.8), Inches(1.0), label, size=14, bold=True, color=INK, align=PP_ALIGN.CENTER)
    set_textbox(
        s,
        Inches(0.8),
        Inches(5.2),
        Inches(11.7),
        Inches(0.9),
        "در هر مرحله زمان‌بندی سرویس‌ها اندازه‌گیری و به کاربر نمایش داده می‌شود تا شفافیت عملکرد سامانه حفظ شود.",
        size=15,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )
    footer(s, 5, total)

    # ----- 6 Features -----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    add_rect(s, 0, 0, prs.slide_width, Inches(1.1), GREEN)
    set_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.5), "قابلیت‌های کلیدی", size=28, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    add_bullets(
        s,
        Inches(0.8),
        Inches(1.5),
        Inches(11.5),
        Inches(5.2),
        [
            "پشتیبانی دوزبانه فارسی و انگلیسی در سؤال و پاسخ.",
            "انتخاب مدل AvalAI یا Ollama و ذخیره‌سازی ترجیح هر کاربر.",
            "تاریخچه گفتگو برای هر کاربر به‌صورت جداگانه.",
            "اقدامات سریع در داشبورد (فروشندگان برتر، خلاصه امروز و ...).",
            "مدیریت کاربران توسط مدیر: نام کاربری، رمز، موبایل و نام نمایشی.",
            "امنیت ورود با کپچای «من ربات نیستم» و کنترل نشست.",
        ],
        size=17,
    )
    footer(s, 6, total)

    # ----- 7 AI Providers -----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    add_rect(s, 0, 0, prs.slide_width, Inches(1.1), GREEN)
    set_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.5), "مدل‌های هوش مصنوعی", size=28, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    add_round_rect(s, Inches(0.7), Inches(1.6), Inches(5.8), Inches(4.4), GREEN_SOFT)
    add_round_rect(s, Inches(6.8), Inches(1.6), Inches(5.8), Inches(4.4), CREAM)
    set_textbox(s, Inches(1.0), Inches(1.9), Inches(5.2), Inches(0.5), "AvalAI (پیش‌فرض)", size=22, bold=True, color=GREEN_DARK, align=PP_ALIGN.RIGHT)
    add_bullets(
        s,
        Inches(1.0),
        Inches(2.6),
        Inches(5.2),
        Inches(3.0),
        [
            "API ابری سازگار با OpenAI",
            "مناسب استقرار روی هاست و سرور",
            "انتخاب پیش‌فرض سامانه",
        ],
        size=16,
    )
    set_textbox(s, Inches(7.1), Inches(1.9), Inches(5.2), Inches(0.5), "Ollama (محلی)", size=22, bold=True, color=GREEN_DARK, align=PP_ALIGN.RIGHT)
    add_bullets(
        s,
        Inches(7.1),
        Inches(2.6),
        Inches(5.2),
        Inches(3.0),
        [
            "اجرا روی سیستم محلی کاربر",
            "مناسب محیط توسعه و آفلاین",
            "نیازمند سرویس ollama serve",
        ],
        size=16,
    )
    footer(s, 7, total)

    # ----- 8 Security & Users -----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    add_rect(s, 0, 0, prs.slide_width, Inches(1.1), GREEN)
    set_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.5), "امنیت و مدیریت کاربران", size=28, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    add_bullets(
        s,
        Inches(0.8),
        Inches(1.5),
        Inches(11.5),
        Inches(5.2),
        [
            "رمز عبور کاربران به‌صورت هش‌شده (PBKDF2) در پایگاه داده ذخیره می‌شود.",
            "نقش مدیر می‌تواند کاربر جدید اضافه کند، ویرایش کند یا حذف کند.",
            "صفحه مدیریت کاربران فقط برای نقش admin قابل مشاهده است.",
            "لایه sql_safety فقط دستورات SELECT امن را می‌پذیرد و دستورات مخرب را رد می‌کند.",
            "نشست کاربران با SessionMiddleware محافظت می‌شود.",
        ],
        size=17,
    )
    footer(s, 8, total)

    # ----- 9 UI -----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    add_rect(s, 0, 0, prs.slide_width, Inches(1.1), GREEN)
    set_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.5), "تجربه کاربری و داشبورد", size=28, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    add_bullets(
        s,
        Inches(0.8),
        Inches(1.5),
        Inches(11.5),
        Inches(5.2),
        [
            "صفحه معرفی (First) با هویت بصری سبز رایامیت.",
            "صفحه ورود با زمینه Rayamate و کپچای تأیید انسانی.",
            "داشبورد سمت راست: تاریخچه گفتگو، اقدامات سریع و تنظیمات مدل/زبان.",
            "امکان جمع‌کردن پنل کناری برای تمرکز بیشتر روی گفتگو.",
            "نمایش جدول نتایج و نمودار سبک پس از اجرای گزارش.",
        ],
        size=17,
    )
    footer(s, 9, total)

    # ----- 10 Tech stack -----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    add_rect(s, 0, 0, prs.slide_width, Inches(1.1), GREEN)
    set_textbox(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.5), "فناوری‌های استفاده‌شده", size=28, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    techs = [
        ("FastAPI + Uvicorn", "سرویس‌دهی API"),
        ("SQLAlchemy + SQLite", "لایه داده"),
        ("Ollama / AvalAI", "مدل زبانی"),
        ("HTML/CSS/JS", "رابط کاربری"),
        ("Docker / cPanel", "استقرار"),
        ("Session Auth", "ورود کاربران"),
    ]
    for i, (t, d) in enumerate(techs):
        col = i % 3
        row = i // 3
        left = Inches(0.7 + col * 4.15)
        top = Inches(1.7 + row * 2.3)
        add_round_rect(s, left, top, Inches(3.9), Inches(1.9), GREEN_SOFT)
        set_textbox(s, left + Inches(0.2), top + Inches(0.45), Inches(3.5), Inches(0.45), t, size=18, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
        set_textbox(s, left + Inches(0.2), top + Inches(1.0), Inches(3.5), Inches(0.4), d, size=14, color=INK, align=PP_ALIGN.CENTER)
    footer(s, 10, total)

    # ----- 11 Closing -----
    s = prs.slides.add_slide(blank)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, CREAM)
    add_rect(s, 0, 0, Inches(0.35), prs.slide_height, GREEN)
    set_textbox(s, Inches(1.0), Inches(2.0), Inches(11.3), Inches(0.7), "جمع‌بندی", size=32, bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER)
    set_textbox(
        s,
        Inches(1.5),
        Inches(2.9),
        Inches(10.3),
        Inches(1.4),
        "رایامیت پلی میان زبان طبیعی و گزارش‌های تراکنشی است؛\nبا کنترل امنیتی، مدیریت کاربران و انتخاب انعطاف‌پذیر مدل هوش مصنوعی.",
        size=18,
        color=INK,
        align=PP_ALIGN.CENTER,
    )
    set_textbox(
        s,
        Inches(1.5),
        Inches(4.6),
        Inches(10.3),
        Inches(0.6),
        "با تشکر از توجه شما",
        size=22,
        bold=True,
        color=GREEN,
        align=PP_ALIGN.CENTER,
    )
    set_textbox(
        s,
        Inches(1.5),
        Inches(5.4),
        Inches(10.3),
        Inches(0.4),
        "سؤال و پاسخ",
        size=16,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    footer(s, 11, total)

    out = "Rayamate_Presentation.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
