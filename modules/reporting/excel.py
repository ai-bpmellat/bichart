"""Excel export for BI chat results.

Sheet layout (0-indexed as requested):
  0 — Data only (headers + rows)
  1 — Explanation / title context
  2 — Analysis
"""

from __future__ import annotations

from typing import Any, Optional

from modules.reporting.labels import column_label
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F6F54")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Tahoma", size=11)
CELL_FONT = Font(name="Tahoma", size=10)
TITLE_FONT = Font(bold=True, name="Tahoma", size=14, color="15523E")
SECTION_FONT = Font(bold=True, name="Tahoma", size=12, color="1F6F54")
THIN = Border(
    left=Side(style="thin", color="D8D3C5"),
    right=Side(style="thin", color="D8D3C5"),
    top=Side(style="thin", color="D8D3C5"),
    bottom=Side(style="thin", color="D8D3C5"),
)
ALT_FILL = PatternFill("solid", fgColor="F4F7F5")


def _is_number(value: Any) -> bool:
    if isinstance(value, bool) or value is None:
        return False
    if isinstance(value, (int, float)):
        return True
    try:
        float(str(value).replace(",", "").replace(";", "").replace(" ", ""))
        return True
    except (TypeError, ValueError):
        return False


def _to_number(value: Any):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    s = str(value).replace(",", "").replace(";", "").replace(" ", "")
    if "." in s:
        return float(s)
    return int(s)


def _autosize(ws, max_width: int = 48) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        longest = 0
        for cell in ws[letter]:
            if cell.value is None:
                continue
            longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_width, max(12, longest + 2))


def generate_llm_excel(
    output_path: str,
    title: str,
    explanation: str,
    data: Optional[list] = None,
    analysis: Optional[str] = None,
) -> str:
    wb = Workbook()

    # --- Sheet 0: Data only ---
    ws_data = wb.active
    ws_data.title = "Data"
    rows = data or []
    if rows:
        columns = list(rows[0].keys())
        for c_idx, col in enumerate(columns, start=1):
            cell = ws_data.cell(1, c_idx, column_label(col))
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = THIN
        for r_idx, row in enumerate(rows, start=2):
            for c_idx, col in enumerate(columns, start=1):
                raw = row.get(col, "")
                cell = ws_data.cell(r_idx, c_idx)
                if _is_number(raw):
                    cell.value = _to_number(raw)
                    cell.number_format = "#,##0.##"
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.value = "" if raw is None else str(raw)
                    cell.alignment = Alignment(horizontal="right" if _looks_rtl(str(raw)) else "left")
                cell.font = CELL_FONT
                cell.border = THIN
                if r_idx % 2 == 0:
                    cell.fill = ALT_FILL
        ws_data.auto_filter.ref = ws_data.dimensions
        ws_data.freeze_panes = "A2"
    else:
        ws_data.cell(1, 1, "No data").font = CELL_FONT
    _autosize(ws_data)

    # --- Sheet 1: Explanation (middle sheet so Analysis is index 2) ---
    ws_info = wb.create_sheet("Explanation", 1)
    ws_info.cell(1, 1, title or "BI Report").font = TITLE_FONT
    ws_info.cell(3, 1, "Explanation").font = SECTION_FONT
    ws_info.cell(4, 1, explanation or "").font = CELL_FONT
    ws_info.cell(4, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws_info.column_dimensions["A"].width = 100
    ws_info.row_dimensions[4].height = 120

    # --- Sheet 2: Analysis ---
    ws_analysis = wb.create_sheet("Analysis", 2)
    ws_analysis.cell(1, 1, "Analysis").font = SECTION_FONT
    ws_analysis.cell(3, 1, analysis or "").font = CELL_FONT
    ws_analysis.cell(3, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws_analysis.column_dimensions["A"].width = 100
    ws_analysis.row_dimensions[3].height = 200

    wb.save(output_path)
    return output_path


def _looks_rtl(text: str) -> bool:
    return any("\u0600" <= ch <= "\u06FF" for ch in text)
